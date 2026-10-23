import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Deep Slate Tech Theme
    C_BG_DARK = RGBColor(11, 15, 25)        # #0B0F19
    C_CARD_DARK = RGBColor(22, 30, 46)      # #161E2E
    C_CARD_BORDER = RGBColor(45, 55, 72)    # #2D3748
    C_TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC
    C_TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    C_BLUE = RGBColor(56, 189, 248)         # #38BDF8
    C_INDIGO = RGBColor(129, 140, 248)      # #818CF8
    C_EMERALD = RGBColor(52, 211, 153)      # #34D399
    C_AMBER = RGBColor(251, 191, 36)        # #FBBF24
    C_RED = RGBColor(248, 113, 113)         # #F87171

    FONT_TITLE = "Calibri"
    FONT_BODY = "Arial"

    def set_slide_bg(slide, color):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = color
        bg_shape.line.fill.background() # no line
        return bg_shape

    def add_header(slide, title_text, category_text="CHUYÊN ĐỀ PHÂN LOẠI MÃ ĐỘC • UIT"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.name = FONT_BODY
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = C_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=C_CARD_DARK, border_color=C_CARD_BORDER):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, C_BG_DARK)

    # Accent decorative element
    dec = add_card(s1, Inches(0.8), Inches(1.2), Inches(3.2), Inches(0.4), bg_color=RGBColor(30, 41, 59), border_color=C_BLUE)
    tf_dec = dec.text_frame
    tf_dec.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_dec = tf_dec.paragraphs[0]
    p_dec.alignment = PP_ALIGN.CENTER
    p_dec.text = "BÁO CÁO ĐỒ ÁN CAO HỌC CNTT - UIT"
    p_dec.font.name = FONT_BODY
    p_dec.font.size = Pt(11)
    p_dec.font.bold = True
    p_dec.font.color.rgb = C_BLUE

    # Title box
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0
    p1 = tf1.paragraphs[0]
    p1.text = "NGHIÊN CỨU & HIỆN THỰC HÓA\nKHUNG IF-MALEVADE"
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Né tránh các bộ phát hiện mã độc Machine Learning Black-Box bằng kỹ thuật kết hợp MalGAN & Deep Reinforcement Learning"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(18)
    p2.font.color.rgb = C_INDIGO
    p2.space_before = Pt(14)

    # Info card at bottom
    info_card = add_card(s1, Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.8), bg_color=C_CARD_DARK)
    tf_info = info_card.text_frame
    tf_info.margin_left = Inches(0.4)
    tf_info.margin_top = Inches(0.3)
    
    p_i1 = tf_info.paragraphs[0]
    p_i1.text = "• Môn học: Chuyên đề Phân loại Mã độc (Project-02)"
    p_i1.font.name = FONT_BODY
    p_i1.font.size = Pt(14)
    p_i1.font.bold = True
    p_i1.font.color.rgb = C_TEXT_WHITE

    p_i2 = tf_info.add_paragraph()
    p_i2.text = "• Paper gốc: A Deep Reinforcement Learning Framework to Evade Black-Box Machine Learning Based IoT Malware Detectors Using GAN-Generated Influential Features (Arif et al., IEEE Access 2023)"
    p_i2.font.name = FONT_BODY
    p_i2.font.size = Pt(13)
    p_i2.font.color.rgb = C_TEXT_MUTED
    p_i2.space_before = Pt(6)

    p_i3 = tf_info.add_paragraph()
    p_i3.text = "• Mô hình tiếp cận: Đánh giá Đa không gian (Dual-Space: Feature-Space Simulation & Problem-Space Binary Execution)"
    p_i3.font.name = FONT_BODY
    p_i3.font.size = Pt(13)
    p_i3.font.color.rgb = C_EMERALD
    p_i3.space_before = Pt(6)

    s1.notes_slide.notes_text_frame.text = (
        "[0:00 - 0:30] LỜI THOẠI DIỄN GIẢ:\n"
        "Kính thưa quý Thầy Cô trong Hội đồng, hôm nay em xin phép trình bày báo cáo đồ án Chuyên đề Phân loại Mã độc: "
        "'Nghiên cứu và Hiện thực hóa khung IF-MalEvade nhằm né tránh các bộ phát hiện mã độc học máy black-box "
        "bằng kỹ thuật kết hợp MalGAN và Deep Reinforcement Learning'.\n"
        "Đề tài hướng tới việc giải quyết bài toán cốt lõi: Làm sao biến đổi mã độc để qua mặt các mô hình phân loại học máy "
        "nhưng vẫn bảo toàn 100% tính thực thi và hành vi độc hại ban đầu."
    )

    # =========================================================================
    # SLIDE 2: ĐẶT VẤN ĐỀ & THÁCH THỨC
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, C_BG_DARK)
    add_header(s2, "Đặt Vấn Đề & Hai Thách Thức Cốt Lõi Trong Né Bộ Dò Mã Độc")

    # 3 Column Cards
    col_w = Inches(3.64)
    c_gap = Inches(0.39)
    top_pos = Inches(1.6)
    card_h = Inches(4.5)

    # Card 1: Bối cảnh
    c1 = add_card(s2, Inches(0.8), top_pos, col_w, card_h)
    tf = c1.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "01. BỐI CẢNH BÀI TOÁN"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "Sự trỗi dậy của ML Detectors"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(12)

    p = tf.add_paragraph()
    p.text = "• Các bộ phát hiện dựa trên học máy (Random Forest, Gradient Boosting, MLP...) đạt độ chính xác rất cao (>99%) trong phân loại tĩnh file PE.\n\n• Tuy nhiên, các mô hình này dễ bị tấn công đối kháng (Adversarial Attacks) theo cơ chế Black-box (kẻ tấn công chỉ truy vấn nhãn đầu ra)."
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(6)

    # Card 2: Thách thức 1
    c2 = add_card(s2, Inches(0.8) + col_w + c_gap, top_pos, col_w, card_h)
    tf = c2.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "02. THÁCH THỨC 1"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    p = tf.add_paragraph()
    p.text = "Bùng nổ Không gian Hành động"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(12)

    p = tf.add_paragraph()
    p.text = "• Một file PE có hơn 3.500 đặc trưng nhị phân (3.269 Import APIs + 254 tên Section).\n\n• Nếu để tác tử DRL tìm kiếm ngẫu nhiên trong toàn bộ không gian này (như DQEAF), không gian hành động quá lớn khiến mô hình khó hội tụ, tốn nhiều lượt sửa không hiệu quả."
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(6)

    # Card 3: Thách thức 2
    c3 = add_card(s2, Inches(0.8) + (col_w + c_gap)*2, top_pos, col_w, card_h)
    tf = c3.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "03. THÁCH THỨC 2"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_RED

    p = tf.add_paragraph()
    p.text = "Bảo toàn Chức năng Nhị phân PE"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(12)

    p = tf.add_paragraph()
    p.text = "• Nghịch lý mã độc: Sửa file để né detector nhưng KHÔNG ĐƯỢC làm hỏng cấu trúc PE (file không bị crash) và PHẢI GIỮ NGUYÊN hành vi độc hại ban đầu.\n\n• Né trên vector số (Feature-space) là chưa đủ; mẫu phải thực thi được trong môi trường Sandbox động thực tế."
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(6)

    # Bottom banner callout
    bot_callout = add_card(s2, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), bg_color=RGBColor(30, 41, 59), border_color=C_BLUE)
    tf_bot = bot_callout.text_frame
    tf_bot.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "🎯 MỤC TIÊU: Kết hợp MalGAN để thu hẹp không gian hành vi + DRL để tìm chuỗi đột biến ngắn nhất + Cuckoo Sandbox để xác thực toàn vẹn."
    p_bot.font.name = FONT_BODY
    p_bot.font.size = Pt(12)
    p_bot.font.bold = True
    p_bot.font.color.rgb = C_BLUE

    s2.notes_slide.notes_text_frame.text = (
        "[0:30 - 1:30] LỜI THOẠI DIỄN GIẢ:\n"
        "Khi nghiên cứu tấn công đối kháng trên mã độc học máy, chúng ta đối mặt với hai thách thức sống còn:\n"
        "Thứ nhất: Không gian đặc trưng của file thực thi PE là rất lớn—hơn 3.500 chiều. Tác tử học tăng cường thông thường sẽ rơi vào tình trạng 'mò kim đáy bể' nếu không có định hướng.\n"
        "Thứ hai, và là quan trọng nhất: Khác với ảnh số hay văn bản, file mã độc nhị phân có ràng buộc thực thi rất nghiêm ngặt. Nếu chúng ta lật bit tùy tiện, file sẽ bị hỏng header, crash khi nạp hoặc mất luôn payload độc hại ban đầu.\n"
        "Do đó, một mẫu đối kháng chỉ được coi là thành công khi nó đồng thời: qua mặt được bộ phân loại và tiếp tục thực thi nguyên vẹn chức năng trong Sandbox."
    )

    # =========================================================================
    # SLIDE 3: NGUYÊN LÝ KHUNG IF-MALEVADE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, C_BG_DARK)
    add_header(s3, "Khung Phương Pháp IF-MalEvade: Sự Kết Hợp Giữa GAN & DRL")

    # 2 Big Blocks: MalGAN (Left) and DQN (Right)
    b_w = Inches(5.66)
    b_h = Inches(4.5)
    
    # Left Block: MalGAN
    b1 = add_card(s3, Inches(0.8), Inches(1.6), b_w, b_h)
    tf = b1.text_frame
    tf.margin_left = tf.margin_right = Inches(0.35)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "GIAI ĐOẠN A: MalGAN TẠO POOL ĐẶC TRƯNG ẢNH HƯỞNG"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "Khắc phục bùng nổ không gian hành động"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = (
        "• Generator (G): Nhận vector malware m + nhiễu z, sinh vector biến đổi Δ. Sử dụng phép toán logic OR (hard_or): x_adv = m ∨ 1(g ≥ 0.5) nhằm đảm bảo CHỈ THÊM đặc trưng lành tính, không xóa đặc trưng gốc.\n\n"
        "• Discriminator (D): Học phân biệt mẫu lành tính thật và mẫu đối kháng m_adv.\n\n"
        "• Black-box Detector: Đóng vai trò teacher giám sát, ép Generator vừa phải đánh lừa D, vừa ép điểm detector xuống dưới ngưỡng độc lực.\n\n"
        "• Output: Trích xuất TOP ĐẶC TRƯNG CÓ ẢNH HƯỞNG CAO NHẤT (Top-K Influential Features) làm không gian hành động tinh chọn cho DRL."
    )
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(10)

    # Right Block: DQN Agent
    b2 = add_card(s3, Inches(6.86), Inches(1.6), b_w, b_h)
    tf = b2.text_frame
    tf.margin_left = tf.margin_right = Inches(0.35)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "GIAI ĐOẠN B: DEEP Q-NETWORK (DQN) TỐI ƯU HÓA BIẾN ĐỔI"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    p = tf.add_paragraph()
    p.text = "Tìm kiếm chuỗi đột biến tối ưu ít bước nhất"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = (
        "• Không gian Trạng thái (State): Vector đặc trưng hiện tại của file PE.\n\n"
        "• Không gian Hành động (Action): Chỉ bốc các ứng viên từ pool đặc trưng ảnh hưởng do MalGAN lọc ra, thay vì tìm kiếm toàn bộ 3.500 đặc trưng.\n\n"
        "• Hàm phần thưởng định hướng (Reward Function):\n"
        "  - Né thành công (score < 0.8): Thưởng lũy thừa lớn theo số bước ngắn:\n"
        "    R = 20 × 100^(-(turn-1)/MAXTURN)\n"
        "  - Chưa né: R = 0 (khuyến khích chuỗi sửa đổi tối thiểu, tránh làm phình to file).\n\n"
        "• Kỹ thuật: Replay Buffer, Target Network, ε-greedy decay."
    )
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(10)

    # Bottom summary
    bot3 = add_card(s3, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), bg_color=RGBColor(30, 41, 59), border_color=C_EMERALD)
    tf_b3 = bot3.text_frame
    tf_b3.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b3 = tf_b3.paragraphs[0]
    p_b3.text = "💡 TỔNG QUAN LUỒNG: MalGAN đóng vai trò 'Bộ lọc chiến lược' ➔ DQN đóng vai trò 'Người thực thi tối ưu chuỗi hành động'."
    p_b3.font.name = FONT_BODY
    p_b3.font.size = Pt(12)
    p_b3.font.bold = True
    p_b3.font.color.rgb = C_EMERALD

    s3.notes_slide.notes_text_frame.text = (
        "[1:30 - 3:00] LỜI THOẠI DIỄN GIẢ:\n"
        "Để vượt qua hai thách thức trên, IF-MalEvade phân chia bài toán thành hai giai đoạn rõ rệt:\n"
        "Ở Giai đoạn A, mô hình MalGAN được huấn luyện đối kháng. Điểm mấu chốt là Generator bị ràng buộc bởi hàm logic OR—nghĩa là "
        "chỉ được bổ sung các đặc trưng lành tính chứ tuyệt đối không xóa import gốc của malware. Kết quả của MalGAN là một bộ "
        "'Top đặc trưng ảnh hưởng cao nhất'.\n"
        "Ở Giai đoạn B, tác tử Deep Q-Network (DQN) sẽ nhận không gian hành động đã được tinh lọc này. Với hàm thưởng thiết kế "
        "dạng hàm mũ nghịch đảo theo số lượt sửa, DQN bị ép phải tìm ra chuỗi chèn đặc trưng ngắn nhất mà vẫn kéo được điểm phát hiện xuống dưới 0.8."
    )

    # =========================================================================
    # SLIDE 4: THIẾT KẾ THỰC NGHIỆM HỢP NHẤT (DUAL-SPACE)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, C_BG_DARK)
    add_header(s4, "Thiết Kế Thực Nghiệm Hợp Nhất: Đa Không Gian (Dual-Space Evaluation)")

    # 2 Comparison Columns
    col_w4 = Inches(5.66)
    col_h4 = Inches(4.5)

    # Left: Mode 1 - Feature Space
    m1 = add_card(s4, Inches(0.8), Inches(1.6), col_w4, col_h4)
    tf = m1.text_frame
    tf.margin_left = tf.margin_right = Inches(0.35)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "CHẾ ĐỘ 1: FEATURE-SPACE EVALUATION"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_INDIGO

    p = tf.add_paragraph()
    p.text = "Mô phỏng toán học & Kiểm chứng hội tụ thuật toán"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = (
        "• Dữ liệu: 2.000 mẫu BODMAS (1.000 Malware + 1.000 Benign).\n"
        "• Tiền xử lý: SelectKBest chọn lọc 300 đặc trưng nhạy cảm nhất, nhị phân hóa theo ngưỡng median.\n"
        "• Không gian hành động: 50 bit ảnh hưởng cao nhất từ MalGAN.\n"
        "• Cơ chế biến đổi: Lật bit nhị phân trực tiếp trên bộ nhớ (0 ➔ 1).\n"
        "• Môi trường: Black-box PyTorch trên GPU Colab, tốc độ hội tụ nhanh (600 episodes trong vài phút).\n"
        "• Mục đích: Đóng vai trò là phép thử nhanh (Fast Prototyping) để kiểm chứng khả năng đánh lừa bộ phân loại trong điều kiện lý tưởng."
    )
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(12)

    # Right: Mode 2 - Problem Space
    m2 = add_card(s4, Inches(6.86), Inches(1.6), col_w4, col_h4)
    tf = m2.text_frame
    tf.margin_left = tf.margin_right = Inches(0.35)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "CHẾ ĐỘ 2: PROBLEM-SPACE EVALUATION"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    p = tf.add_paragraph()
    p.text = "Thực chiến trên File PE Nhị phân & Sandbox"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = (
        "• Dữ liệu: 3.500 file PE thực tế (2.500 malware RanDS + 1.000 PE Windows lành tính).\n"
        "• Không gian đặc trưng: 3.523 bit đầy đủ (3.269 imports + 254 sections) bám sát paper.\n"
        "• Cơ chế biến đổi: 7 actions can thiệp nhị phân thực tế bằng thư viện LIEF, pefile, UPX.\n"
        "• Ràng buộc nghiêm ngặt: Xuất file .exe đối kháng hoàn chỉnh.\n"
        "• Kiểm chứng động Cuckoo Sandbox v3: Đối chiếu toàn bộ trường signatures hành vi động (so sánh bản gốc vs mutant).\n"
        "• Kiểm chứng bên ngoài: Quét VirusTotal thực tế trên 50+ Antivirus engines."
    )
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(12)

    # Bottom Callout
    bot4 = add_card(s4, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), bg_color=RGBColor(30, 41, 59), border_color=C_BLUE)
    tf_b4 = bot4.text_frame
    tf_b4.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b4 = tf_b4.paragraphs[0]
    p_b4.text = "🔄 KIẾN TRÚC MÔ-ĐUN: Đồ án tích hợp cả hai chế độ qua cờ tham số (--mode), tạo thành quy trình hoàn chỉnh từ lý thuyết đến thực tiễn."
    p_b4.font.name = FONT_BODY
    p_b4.font.size = Pt(12)
    p_b4.font.bold = True
    p_b4.font.color.rgb = C_BLUE

    s4.notes_slide.notes_text_frame.text = (
        "[3:00 - 4:30] LỜI THOẠI DIỄN GIẢ:\n"
        "Để đánh giá toàn diện, em đã hợp nhất thực nghiệm thành một kiến trúc Đa không gian gồm 2 chế độ:\n"
        "Chế độ 1 là Feature-Space Evaluation: Em thực hiện trên tập BODMAS với 300 đặc trưng chọn lọc. Ở đây, tác tử lật bit trực tiếp trên vector số. "
        "Mục đích của chế độ này là đóng vai trò 'phép thử nhanh', kiểm chứng xem thuật toán GAN và DQN có thể hội tụ tối ưu hay không.\n"
        "Chế độ 2 là Problem-Space Evaluation: Sau khi thuật toán đã chứng minh tính hiệu quả, em đưa trực tiếp vào file PE nhị phân thật với 3.523 chiều đặc trưng, "
        "thao tác bằng các công cụ biến đổi binary chuyên sâu và kiểm tra nghiêm ngặt qua Cuckoo Sandbox cũng như VirusTotal."
    )

    # =========================================================================
    # SLIDE 5: KẾT QUẢ THỰC NGHIỆM HỢP NHẤT
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, C_BG_DARK)
    add_header(s5, "Bảng Kết Quả Thực Nghiệm Đối Sánh Hợp Nhất (Table 4)")

    # Data Table
    table_shape = s5.shapes.add_table(6, 6, Inches(0.8), Inches(1.5), Inches(11.7), Inches(3.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(1.8)
    table.columns[3].width = Inches(1.8)
    table.columns[4].width = Inches(1.8)
    table.columns[5].width = Inches(1.8)

    headers = ["Mô hình mục tiêu", "AUC (%)", "ODR Gốc (%)", "Paper ADR (%)", "Chế độ 1 ADR (%)", "Chế độ 2 ADR (%)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 41, 59)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_BODY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_BLUE
        p.alignment = PP_ALIGN.CENTER

    data_rows = [
        ["Decision Tree", "94.14 / 99.82", "94.50 / 100.0", "8.10%", "20.50%", "76.00%"],
        ["Gradient Boosting", "99.58 / 99.63", "92.50 / 100.0", "13.56%", "20.50%", "96.00%"],
        ["Multi-Layer Perceptron", "99.45 / 99.81", "96.50 / 98.0", "7.69%", "45.00%", "72.00%"],
        ["Random Forest", "99.84 / 99.72", "85.50 / 98.0", "9.10%", "16.50%", "60.00%"],
        ["TRUNG BÌNH ADR", "—", "—", "9.61%", "25.63%", "76.00%"]
    ]

    for row_idx, row in enumerate(data_rows):
        is_summary = (row_idx == len(data_rows) - 1)
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            if is_summary:
                cell.fill.fore_color.rgb = RGBColor(30, 41, 59)
            else:
                cell.fill.fore_color.rgb = C_CARD_DARK if row_idx % 2 == 0 else RGBColor(17, 24, 39)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = FONT_BODY
            p.font.size = Pt(12)
            p.font.bold = is_summary or (col_idx in [3, 4, 5])
            if col_idx == 3:
                p.font.color.rgb = C_BLUE
            elif col_idx == 4:
                p.font.color.rgb = C_INDIGO
            elif col_idx == 5:
                p.font.color.rgb = C_AMBER if not is_summary else C_AMBER
            else:
                p.font.color.rgb = C_TEXT_WHITE
            p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

    # 3 Summary Highlight Cards at bottom
    c_w5 = Inches(3.64)
    c_gap5 = Inches(0.39)
    top_pos5 = Inches(5.0)
    card_h5 = Inches(2.0)

    # Highlight 1: RF Best
    h1 = add_card(s5, Inches(0.8), top_pos5, c_w5, card_h5)
    tf = h1.text_frame
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "🎯 RANDOM FOREST NÉ TỐT NHẤT"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    p = tf.add_paragraph()
    p.text = "• Chế độ 1: ADR chỉ 16.50% (tụt 69 điểm).\n• Chế độ 2: ADR đạt 60.00% (tụt 38 điểm).\n➔ Điểm né tốt nhất trong cả hai môi trường."
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(4)

    # Highlight 2: Cuckoo Sandbox
    h2 = add_card(s5, Inches(0.8) + c_w5 + c_gap5, top_pos5, c_w5, card_h5)
    tf = h2.text_frame
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "🛡️ CUCKOO SANDBOX V3"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "• 50/50 cặp đối kháng chạy dynamic.\n• 42/50 cặp (84%) khớp 100% signatures.\n➔ Chứng minh file biến đổi bảo toàn nguyên vẹn chức năng độc hại ban đầu."
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(4)

    # Highlight 3: VirusTotal
    h3 = add_card(s5, Inches(0.8) + (c_w5 + c_gap5)*2, top_pos5, c_w5, card_h5)
    tf = h3.text_frame
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "🌐 QUÉT VIRUSTOTAL THỰC TẾ"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    p = tf.add_paragraph()
    p.text = "• Quét trên 50+ Antivirus engines.\n• Tỷ lệ phát hiện gốc: 76.8% ➔ Sau sửa: 62.4%.\n• Tỷ lệ né trung bình: 20.5% (cao nhất 75%).\n➔ Đánh bại một số AV thương mại."
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(4)

    s5.notes_slide.notes_text_frame.text = (
        "[4:30 - 7:00] LỜI THOẠI DIỄN GIẢ:\n"
        "Kính thưa Hội đồng, đây là bảng số liệu tổng hợp đối sánh giữa Paper gốc và hai chế độ thực nghiệm của em:\n"
        "Tại Chế độ 1 (Vector BODMAS): Thuật toán đạt kết quả rất ấn tượng, ADR trung bình giảm xuống còn 25,63%. Đặc biệt trên Random Forest, "
        "tỷ lệ bị bắt chỉ còn 16,5%, chứng minh chính sách của DQN hoạt động rất sắc bén khi tối ưu trên vector số.\n"
        "Tại Chế độ 2 (File PE thực tế): Khi đưa vào môi trường nhị phân thật, tỷ lệ bị bắt trung bình là 76%. Trong đó Random Forest vẫn là mô hình "
        "bị né tốt nhất với ADR 60%. Điều quan trọng nhất ở chế độ này là kết quả kiểm chứng động: 42 trên 50 mẫu (tức 84%) giữ nguyên 100% "
        "hành vi độc lực trên Cuckoo Sandbox, và khi quét trên VirusTotal thực tế, mẫu đối kháng đã né được trung bình 20,5% các engine quét."
    )

    # =========================================================================
    # SLIDE 6: THẢO LUẬN CHUYÊN SÂU
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, C_BG_DARK)
    add_header(s6, "Thảo Luận Chuyên Sâu: Ranh Giới Giữa Vector & File PE Thực Tế")

    # 3 Analytical Points
    col_w6 = Inches(3.64)
    c_gap6 = Inches(0.39)
    top_pos6 = Inches(1.6)
    card_h6 = Inches(4.5)

    # Card 1: Inverse Mapping
    a1 = add_card(s6, Inches(0.8), top_pos6, col_w6, card_h6)
    tf = a1.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "01. NGHỊCH LÝ ÁNH XẠ NGƯỢC"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "Inverse Feature Mapping Problem"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = (
        "• Trên không gian vector (Chế độ 1), việc lật bit 0 ➔ 1 chỉ là phép toán trong bộ nhớ.\n\n"
        "• Nhưng trong thực tế (Chế độ 2), việc thêm 1 API vào Import Table đòi hỏi tái cấu trúc toàn bộ PE Header, căn chỉnh Virtual Address và File Offset.\n\n"
        "• Nhiều biến đổi vector không thể thực thi được trên file nhị phân nếu không làm hỏng tính toàn vẹn của mã máy."
    )
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(8)

    # Card 2: Sandbox Filter
    a2 = add_card(s6, Inches(0.8) + col_w6 + c_gap6, top_pos6, col_w6, card_h6)
    tf = a2.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "02. RÀO CHẮN SANDBOX ĐỘNG"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    p = tf.add_paragraph()
    p.text = "Bộ lọc tính hợp lệ khắt khe"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = (
        "• Ở Chế độ 1, mọi mẫu đối kháng đều được chấp nhận miễn là xác suất < 0.8.\n\n"
        "• Ở Chế độ 2, cổng Cuckoo Sandbox v3 loại bỏ mọi mẫu làm crash ứng dụng hoặc làm mất signatures hành vi gốc.\n\n"
        "• Ràng buộc thực tế này đã thu hẹp đáng kể không gian biến đổi, lý giải vì sao ADR ở file PE thật (76%) cao hơn so với mô phỏng vector (25.6%)."
    )
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(8)

    # Card 3: Feature Space Sparsity
    a3 = add_card(s6, Inches(0.8) + (col_w6 + c_gap6)*2, top_pos6, col_w6, card_h6)
    tf = a3.text_frame
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.35)
    p = tf.paragraphs[0]
    p.text = "03. ĐỘ THƯA CỦA KHÔNG GIAN"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    p = tf.add_paragraph()
    p.text = "Tính chống chịu của Mô hình thật"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = (
        "• Chế độ 1 dùng 300 đặc trưng chọn lọc qua ANOVA F-value ➔ Các đặc trưng này có tính phân tách rất nhạy cảm, lật vài bit là mô hình gãy nhánh phân loại.\n\n"
        "• Chế độ 2 dùng 3.523 đặc trưng thưa ➔ Các mô hình như Gradient Boosting học được các cây quyết định rất sâu, việc chèn thêm một vài imports lành tính không đủ làm đảo ngược phán quyết phân loại."
    )
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_before = Pt(8)

    # Bottom summary
    bot6 = add_card(s6, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), bg_color=RGBColor(30, 41, 59), border_color=C_INDIGO)
    tf_b6 = bot6.text_frame
    tf_b6.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b6 = tf_b6.paragraphs[0]
    p_b6.text = "🧠 BÀI HỌC HỌC THUẬT: Đồ án chỉ ra rằng các nghiên cứu Adversarial ML trên vector cần được kiểm chứng trên Problem-Space để tránh ảo tưởng về hiệu năng né tránh."
    p_b6.font.name = FONT_BODY
    p_b6.font.size = Pt(12)
    p_b6.font.bold = True
    p_b6.font.color.rgb = C_INDIGO

    s6.notes_slide.notes_text_frame.text = (
        "[7:00 - 8:30] LỜI THOẠI DIỄN GIẢ:\n"
        "Sự chênh lệch giữa tỷ lệ né 74,4% ở Chế độ 1 và 24% ở Chế độ 2 chính là đóng góp học thuật sâu sắc nhất của đồ án:\n"
        "Thứ nhất, bài toán Inverse Feature Mapping: Trong không gian vector, lật bit rất dễ, nhưng khi chuyển thành byte thực tế, "
        "việc sửa đổi cấu trúc nhị phân bị hạn chế bởi sự toàn vẹn của PE Header.\n"
        "Thứ hai, vai trò của Sandbox: Khi có cổng kiểm soát hành vi động, chúng ta buộc phải loại bỏ những đột biến làm chết file. "
        "Chính sự nghiêm ngặt này khiến tỷ lệ né thực tế không thể dễ dãi như trên mô phỏng.\n"
        "Thứ ba, không gian 3.523 đặc trưng thực tế có tính chống chịu đối kháng cao hơn nhiều so với không gian 300 đặc trưng rút gọn.\n"
        "Đây là phát hiện rất quan trọng, cảnh báo rằng nếu chỉ nghiên cứu tấn công trên mảng số, chúng ta sẽ có một 'ảo tưởng' về khả năng né tránh."
    )

    # =========================================================================
    # SLIDE 7: KẾT LUẬN & ĐÓNG GÓP
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, C_BG_DARK)
    add_header(s7, "Kết Luận & Giá Trị Đóng Góp Của Đề Tài")

    # Main Card
    main_c7 = add_card(s7, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.5))
    tf = main_c7.text_frame
    tf.margin_left = tf.margin_right = Inches(0.4)
    tf.margin_top = Inches(0.35)

    p = tf.paragraphs[0]
    p.text = "CÁC KẾT QUẢ ĐÃ ĐẠT ĐƯỢC:"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    points = [
        ("1. Hiện thực hóa thành công khung IF-MalEvade:", "Xây dựng hoàn chỉnh kiến trúc kết hợp MalGAN (lọc đặc trưng ảnh hưởng) và Deep Q-Network (tối ưu hóa chuỗi biến đổi)."),
        ("2. Xây dựng Khung thực nghiệm Đa không gian:", "Tích hợp cả 2 chế độ: Feature-space (mô phỏng nhanh trên Colab) và Problem-space (thực thi file nhị phân trên Lab) trong một kiến trúc thống nhất."),
        ("3. Đạt kết quả thực nghiệm toàn diện:", "Chứng minh tính khả thi của thuật toán với ADR 25.6% trên vector; thực hiện thành công biến đổi file PE thật với 84% mẫu bảo toàn hành vi trên Cuckoo Sandbox và né 20.5% trên VirusTotal."),
        ("4. Đóng góp bài học phản biện khoa học:", "Chỉ rõ ranh giới và khoảng cách giữa nghiên cứu Adversarial ML thuần lý thuyết và bài toán an ninh mã độc thực chiến trong thế giới thực.")
    ]

    for title, desc in points:
        p = tf.add_paragraph()
        p.text = f"{title} "
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_WHITE
        p.space_before = Pt(10)

        run = p.add_run()
        run.text = desc
        run.font.name = FONT_BODY
        run.font.size = Pt(13)
        run.font.bold = False
        run.font.color.rgb = C_TEXT_MUTED

    # Bottom Contact / Q&A Card
    bot7 = add_card(s7, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), bg_color=RGBColor(30, 41, 59), border_color=C_EMERALD)
    tf_b7 = bot7.text_frame
    tf_b7.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b7 = tf_b7.paragraphs[0]
    p_b7.text = "🙏 XIN CHÂN THÀNH CẢM ƠN QUÝ THẦY CÔ TRONG HỘI ĐỒNG! (SẴN SÀNG CHO PHẦN HỎI ĐÁP Q&A)"
    p_b7.font.name = FONT_BODY
    p_b7.font.size = Pt(13)
    p_b7.font.bold = True
    p_b7.font.color.rgb = C_EMERALD
    p_b7.alignment = PP_ALIGN.CENTER

    s7.notes_slide.notes_text_frame.text = (
        "[8:30 - 10:00] LỜI THOẠI DIỄN GIẢ:\n"
        "Tóm lại, đồ án đã hoàn thành trọn vẹn các mục tiêu đề ra: từ việc hiện thực hóa khung IF-MalEvade, xây dựng quy trình kiểm thử "
        "đa không gian, đến việc chứng minh được tính toàn vẹn chức năng của mã độc qua Sandbox thực tế.\n"
        "Em xin chân thành cảm ơn quý Thầy Cô đã lắng nghe và em rất mong nhận được những câu hỏi góp ý quý báu từ Hội đồng!"
    )

    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "BaoCao_IF-MalEvade_10Phut.pptx"
    create_deck(out)
